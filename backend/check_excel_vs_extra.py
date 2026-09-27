import openpyxl

wb = openpyxl.load_workbook("C:/coalnexus/data/indian_coal_mines.xlsx")
ws = wb.active

print("Checking if Umrer, Majri, or Ballarpur are in the excel spreadsheet...")
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if row and row[3]:
        name = str(row[3]).lower()
        if "umrer" in name or "majri" in name or "ballarpur" in name:
            print(f"Found in Excel: {row[0]}, {row[1]}, {row[2]}, {row[3]}")
print("Finished check.")
