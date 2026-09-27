import os
import sys
import openpyxl
from sqlalchemy.orm import Session
from geoalchemy2.elements import WKTElement

# Add project root to python path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.database import SessionLocal
from app.models.models import Mine, Profile, UserMineAssignment, MineSensor, RiskScore, Inspection, Violation, Alert

def import_master_data():
    db: Session = SessionLocal()
    try:
        file_path = "C:/coalnexus/data/indian_coal_mines.xlsx"
        print(f"Loading Excel file from {file_path}...")
        wb = openpyxl.load_workbook(file_path)
        # Select active sheet or by name if active sheet isn't named right
        ws = wb.active
        
        print("Parsing rows and inserting/updating mines...")
        
        rows_read = 0
        valid_rows = 0
        inserted_count = 0
        updated_count = 0
        skipped_count = 0
        rejected_count = 0
        
        # Row 1 is headers. Let's iterate from row 2
        for idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), start=2):
            if not row:
                continue
            
            rows_read += 1
            sl_no = row[0]
            if sl_no is None:
                skipped_count += 1
                continue
                
            state = str(row[1]).strip() if row[1] else None
            district = str(row[2]).strip() if row[2] else None
            mine_name = str(row[3]).strip() if row[3] else None
            
            if not mine_name:
                rejected_count += 1
                print(f"Row {idx}: Rejected due to missing Mine Name.")
                continue
                
            try:
                production = float(row[4]) if row[4] is not None else 0.0
            except (ValueError, TypeError):
                production = 0.0
                
            owner_code = str(row[5]).strip() if row[5] else None
            owner_name = str(row[6]).strip() if row[6] else None
            commodity = str(row[7]).strip() if row[7] else None
            ownership_type = str(row[8]).strip() if row[8] else None
            mine_type = str(row[9]).strip() if row[9] else None
            
            try:
                lat = float(row[10]) if row[10] is not None else 0.0
                lon = float(row[11]) if row[11] is not None else 0.0
            except (ValueError, TypeError):
                lat, lon = 0.0, 0.0
                
            source = str(row[12]).strip() if row[12] else None
            coordinate_accuracy = str(row[13]).strip() if row[13] else None
            
            valid_rows += 1
            
            # Generate deterministic mine_code from owner_code and sl_no to ensure idempotency
            mine_code = f"IND_{owner_code or 'UNKNOWN'}_{sl_no}".upper()
            
            loc_wkt = f"POINT({lon} {lat})"
            
            # Check if mine already exists by mine_code
            existing_mine = db.query(Mine).filter(Mine.mine_code == mine_code).first()
            
            if existing_mine:
                # Update fields
                existing_mine.name = mine_name
                existing_mine.company = owner_name # matching previous convention or field mapping
                existing_mine.district = district
                existing_mine.state = state
                existing_mine.latitude = lat
                existing_mine.longitude = lon
                existing_mine.location = WKTElement(loc_wkt, srid=4326)
                existing_mine.owner_code = owner_code
                existing_mine.owner_name = owner_name
                existing_mine.ownership_type = ownership_type
                existing_mine.commodity = commodity
                existing_mine.mine_type = mine_type
                existing_mine.production_hist = production
                existing_mine.coordinate_accuracy = coordinate_accuracy
                existing_mine.source = source
                updated_count += 1
            else:
                # Insert new mine
                mine_obj = Mine(
                    id=f"srv_ind_{sl_no}",
                    local_id=f"loc_mine_{sl_no}",
                    name=mine_name,
                    mine_code=mine_code,
                    company=owner_name,
                    district=district,
                    state=state,
                    latitude=lat,
                    longitude=lon,
                    location=WKTElement(loc_wkt, srid=4326),
                    status="active",
                    owner_code=owner_code,
                    owner_name=owner_name,
                    ownership_type=ownership_type,
                    commodity=commodity,
                    mine_type=mine_type,
                    production_hist=production,
                    coordinate_accuracy=coordinate_accuracy,
                    source=source,
                    local_version=1
                )
                db.add(mine_obj)
                inserted_count += 1
                
        db.commit()
        
        # Get final count
        final_count = db.query(Mine).count()
        
        print("\n==================================================")
        print("IMPORT REPORT")
        print("==================================================")
        print(f"Excel rows read: {rows_read}")
        print(f"Valid rows: {valid_rows}")
        print(f"Inserted: {inserted_count}")
        print(f"Updated: {updated_count}")
        print(f"Skipped (empty sl_no): {skipped_count}")
        print(f"Rejected: {rejected_count}")
        print(f"Final database mine count: {final_count}")
        print("==================================================")
        
        # Re-assign demo profiles to a couple of imported mines so testing works smoothly without integrity failures
        print("Checking demo profile assignments...")
        all_mines = db.query(Mine).limit(5).all()
        if len(all_mines) >= 3:
            officer = db.query(Profile).filter(Profile.id == "demo-officer-id-222").first()
            if officer:
                # Clear existing assignments first to avoid duplicates
                db.query(UserMineAssignment).filter(UserMineAssignment.profile_id == officer.id).delete()
                db.add(UserMineAssignment(profile_id=officer.id, mine_id=all_mines[0].id))
                db.add(UserMineAssignment(profile_id=officer.id, mine_id=all_mines[1].id))
            
            inspector = db.query(Profile).filter(Profile.id == "demo-inspector-id-333").first()
            if inspector:
                db.query(UserMineAssignment).filter(UserMineAssignment.profile_id == inspector.id).delete()
                db.add(UserMineAssignment(profile_id=inspector.id, mine_id=all_mines[1].id))
                db.add(UserMineAssignment(profile_id=inspector.id, mine_id=all_mines[2].id))
            db.commit()
            print("Demo roles re-assigned successfully.")
            
    except Exception as e:
        db.rollback()
        print(f"Error during import: {e}")
        raise e
    finally:
        db.close()

if __name__ == "__main__":
    import_master_data()
